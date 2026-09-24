import io
import re
from typing import List, Optional, TypedDict


class ParsedSection(TypedDict):
    section_number: Optional[str]
    heading: Optional[str]
    page_number: Optional[int]
    content: str


class ParseResult(TypedDict):
    total_chars: int
    sections: List[ParsedSection]


_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,4})\s+(.{1,200}?)\s*$")


def _push(sections: List[ParsedSection], buffer: List[str], current: dict) -> None:
    content = "\n".join(buffer).strip()
    if not content:
        return
    sections.append(
        {
            "section_number": current.get("section_number"),
            "heading": current.get("heading"),
            "page_number": current.get("page_number"),
            "content": content,
        }
    )


def parse_pdf(data: bytes) -> ParseResult:
    import pdfplumber

    sections: List[ParsedSection] = []
    current = {"section_number": None, "heading": None, "page_number": 1}
    buffer: List[str] = []
    total_chars = 0

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            total_chars += len(text)
            for raw in text.splitlines():
                line = raw.rstrip()
                if not line.strip():
                    buffer.append("")
                    continue
                m = _HEADING_RE.match(line)
                if m:
                    _push(sections, buffer, current)
                    buffer = []
                    current = {
                        "section_number": m.group(1),
                        "heading": m.group(2).strip(),
                        "page_number": page_idx,
                    }
                else:
                    if current.get("page_number") is None:
                        current["page_number"] = page_idx
                    buffer.append(line)

    _push(sections, buffer, current)

    if not sections and total_chars == 0:
        raise ValueError("No text could be extracted from PDF")

    return {"total_chars": total_chars, "sections": sections}


def parse_docx(data: bytes) -> ParseResult:
    from docx import Document as DocxDocument

    sections: List[ParsedSection] = []
    current = {"section_number": None, "heading": None, "page_number": None}
    buffer: List[str] = []
    total_chars = 0

    doc = DocxDocument(io.BytesIO(data))

    for para in doc.paragraphs:
        text = (para.text or "").rstrip()
        style = (para.style.name or "") if para.style else ""
        is_heading = style.lower().startswith("heading")

        if is_heading and text:
            _push(sections, buffer, current)
            buffer = []
            m = _HEADING_RE.match(text)
            if m:
                current = {
                    "section_number": m.group(1),
                    "heading": m.group(2).strip(),
                    "page_number": None,
                }
            else:
                current = {
                    "section_number": None,
                    "heading": text,
                    "page_number": None,
                }
            continue

        if not text:
            buffer.append("")
            continue

        m = _HEADING_RE.match(text)
        if m and len(text) <= 200:
            _push(sections, buffer, current)
            buffer = []
            current = {
                "section_number": m.group(1),
                "heading": m.group(2).strip(),
                "page_number": None,
            }
        else:
            buffer.append(text)
            total_chars += len(text)

    _push(sections, buffer, current)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join((c.text or "").strip() for c in row.cells)
            if row_text.strip():
                total_chars += len(row_text)
                sections.append(
                    {
                        "section_number": None,
                        "heading": "Table row",
                        "page_number": None,
                        "content": row_text,
                    }
                )

    if not sections and total_chars == 0:
        raise ValueError("No text could be extracted from DOCX")

    return {"total_chars": total_chars, "sections": sections}


def parse(mime_type: str, data: bytes) -> ParseResult:
    if mime_type == "application/pdf":
        return parse_pdf(data)
    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return parse_docx(data)
    raise ValueError(f"Unsupported mime type: {mime_type}")
